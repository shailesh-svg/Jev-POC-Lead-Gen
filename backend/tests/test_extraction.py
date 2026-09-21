"""Documents a sales team actually drops on the lead tab."""
import io
import pytest
from docx import Document
from openpyxl import Workbook
from backend.extraction import extract_document
from fastapi import HTTPException
from test_app import client, pdf_bytes

def docx_bytes(paragraphs=(), rows=()):
    document = Document()
    for text in paragraphs:
        document.add_paragraph(text)
    if rows:
        table = document.add_table(rows=len(rows), cols=len(rows[0]))
        for r, row in enumerate(rows):
            for c, value in enumerate(row):
                table.cell(r, c).text = value
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()

def xlsx_bytes(rows):
    workbook = Workbook()
    sheet = workbook.active
    for row in rows:
        sheet.append(row)
    out = io.BytesIO()
    workbook.save(out)
    return out.getvalue()

def test_docx_paragraphs_and_tables_come_through():
    text, kind = extract_document(
        docx_bytes(
            paragraphs=['Freed, ambient scribe for community clinics, 26,000 clinicians.'],
            rows=[['Company', 'Contact'], ['Penciled', 'Shawn Shivdat']],
        ),
        'leads.docx',
    )
    assert kind == 'docx'
    assert 'ambient scribe' in text
    assert 'Penciled\tShawn Shivdat' in text   # rows stay tab separated, so each is one lead

def test_xlsx_rows_become_tab_separated_lines():
    text, kind = extract_document(
        xlsx_bytes([
            ['company', 'contact', 'notes'],
            ['Freed', 'Erez Druk', 'Ambient scribe, 26k clinicians'],
            ['Penciled', 'Shawn Shivdat', 'PT front office, 50 clinics'],
        ]),
        'leads.xlsx',
    )
    assert kind == 'xlsx'
    lines = text.splitlines()
    assert len(lines) == 3
    assert lines[1] == 'Freed\tErez Druk\tAmbient scribe, 26k clinicians'

def test_blank_cells_and_rows_do_not_break_a_sheet():
    text, _ = extract_document(
        xlsx_bytes([
            ['company', 'contact'],
            ['Freed', None],
            [None, None],
            ['Penciled', 'Shawn Shivdat, founder and chief executive'],
        ]),
        'leads.xlsx',
    )
    assert len(text.splitlines()) == 3   # the empty row is dropped

def test_pdf_still_works_and_plain_text_passes_through():
    text, kind = extract_document(pdf_bytes('Freed is an ambient scribe used by 26,000 clinicians.'), 'lead.pdf')
    assert kind == 'pdf' and 'ambient scribe' in text
    csv = b'company,contact\nFreed,Erez Druk\nPenciled,Shawn Shivdat\n'
    text, kind = extract_document(csv, 'leads.csv')
    assert kind == 'csv' and text.startswith('company,contact')

def test_legacy_office_formats_say_what_to_do():
    for name, expected in (('leads.doc', 'Word 97-2003'), ('leads.xls', 'Excel 97-2003')):
        with pytest.raises(HTTPException) as caught:
            extract_document(b'anything', name)
        assert caught.value.status_code == 415
        assert expected in caught.value.detail
        assert 'Save it as' in caught.value.detail

def test_unsupported_and_unreadable_files_fail_clearly():
    with pytest.raises(HTTPException) as caught:
        extract_document(b'data', 'leads.png')
    assert caught.value.status_code == 415
    with pytest.raises(HTTPException) as caught:
        extract_document(b'not really a workbook', 'leads.xlsx')
    assert caught.value.status_code == 422
    with pytest.raises(HTTPException) as caught:
        extract_document(b'tiny', 'leads.txt')
    assert caught.value.status_code == 422   # nothing worth scoring

def test_lead_extract_endpoint_accepts_a_spreadsheet(client):
    rows = xlsx_bytes([['company', 'notes'], ['Freed', 'Ambient scribe for community clinics']])
    r = client.post('/api/lead-extract', files={'file': ('leads.xlsx', rows)})
    assert r.status_code == 200, r.text
    body = r.json()
    assert body['kind'] == 'xlsx'
    assert body['filename'] == 'leads.xlsx'
    assert body['characters'] == len(body['text'])
    assert 'Freed\tAmbient scribe' in body['text']

def test_lead_extract_endpoint_rejects_the_wrong_thing(client):
    assert client.post('/api/lead-extract', files={'file': ('leads.png', b'no')}).status_code == 415
    assert client.post('/api/lead-extract', files={'file': ('big.txt', b'x' * (10 * 1024 * 1024 + 1))}).status_code == 413

def test_classification_still_takes_pdf_only(client):
    """Widening the lead tab must not widen the PDF classifier."""
    r = client.post('/api/classification', files={'file': ('leads.xlsx', xlsx_bytes([['a', 'b']]))})
    assert r.status_code in (409, 415)
